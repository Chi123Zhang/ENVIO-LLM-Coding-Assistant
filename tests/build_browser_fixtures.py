"""Generate ONLY synthetic software fixtures for manual browser acceptance.

Run from the repository: python tests/build_browser_fixtures.py /tmp/orcca-fixtures
No study files, credentials, or network access are used.
"""
import calendar
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook
from orcca_audit.demo import demo_data, public_documents


def build(target):
    target = Path(target)
    target.mkdir(parents=True, exist_ok=True)
    data = demo_data()
    for role in ("panel", "monthly", "annual", "tracker"):
        data[role].drop(columns=["source_file", "_source_row"]).to_csv(target / f"{role}.csv", index=False)
    changed = data["panel"].drop(columns=["source_file", "_source_row"]).copy()
    changed.loc[0, ["Deaths", "y_lo", "y_hi"]] = 999
    changed.to_csv(target / "panel_changed.csv", index=False)
    (target / "specs.json").write_text(json.dumps(data["specifications"]), encoding="utf-8")
    (target / "evidence.txt").write_text(public_documents()[0]["text"], encoding="utf-8")
    (target / "evidence_changed.txt").write_text("SYNTHETIC replacement evidence. Zebra marker only.\n", encoding="utf-8")
    (target / "bad_specs.json").write_text('{"invalid": true}', encoding="utf-8")
    rows = []
    definitions = []
    fields = []
    for base, suffix in [("hcs_orc209", "l"), ("hcs_orc210", "m")]:
        for slot in range(1, 31):
            for field in (f"{base}{suffix}_{slot}", f"{base}v_{slot}"):
                fields.append(field)
                definitions.append(["SYNTHETIC_VERSION", field])
    headers = ["hcs_orc2record_id", "redcap_event_name", "hcs_orc2version", "hcs_orc2_showorc09", "hcs_orc2_showorc10"] + fields
    for county in ("DEMO_A", "DEMO_B", "DEMO_C"):
        for month in range(1, 9):
            row = [county, f"{calendar.month_name[month].lower()}_2024_arm_1", "SYNTHETIC_VERSION", 1, 1]
            row += [1 if f in ("hcs_orc209l_1", "hcs_orc210m_1") else 0 if f in ("hcs_orc209v_1", "hcs_orc210v_1") else None for f in fields]
            rows.append(row)
    wb = Workbook()
    ws = wb.active
    ws.title = "orcca_tracker_v2"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    mapping = wb.create_sheet("Sheet1")
    mapping.append(["id", "county"])
    for county in ("DEMO_A", "DEMO_B", "DEMO_C"):
        mapping.append([county, county])
    wb.save(target / "synthetic_tracker.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "ORCCAT_DD"
    ws.append(["Version", "Variable / Field Name"])
    for row in definitions:
        ws.append(row)
    wb.save(target / "synthetic_dictionary.xlsx")


if __name__ == "__main__":
    build(sys.argv[1])
