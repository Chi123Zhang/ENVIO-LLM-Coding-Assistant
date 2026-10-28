"""Bounded, in-memory parsers. File extensions are not used as format evidence."""
from io import BytesIO, StringIO
import csv
import hashlib
import zipfile

import pandas as pd

from .engine import AuditError

MAX_BYTES = 20 * 1024 * 1024


def read_workbook(data, positional_sheets=()):
    if len(data) > MAX_BYTES:
        raise AuditError("File exceeds the 20 MB audit limit.")
    try:
        with zipfile.ZipFile(BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                raise AuditError("Workbook expansion exceeds the 100 MB limit.")
            if "xl/workbook.xml" not in z.namelist() or any("vbaProject" in n for n in z.namelist()):
                raise AuditError("Expected a macro-free OOXML workbook.")
        import openpyxl
        book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False)
        sheets = {}
        for ws in book:
            if ws.max_row * ws.max_column > 2_000_000:
                raise AuditError("Worksheet dimensions exceed the audit limit.")
            rows = list(ws.values)
            if not rows:
                continue
            nonempty_headers = [v for v in rows[0] if v is not None]
            duplicates = len(set(nonempty_headers)) != len(nonempty_headers)
            if duplicates and ws.title not in positional_sheets:
                raise AuditError("Duplicate worksheet headers require explicit review.")
            if duplicates:
                # Explicitly approved positional lookup sheets retain coordinates;
                # duplicate labels never become a semantic field mapping.
                names = [f"column_{i + 1}" for i in range(len(rows[0]))]
            elif any(v is None for v in rows[0]):
                # Auxiliary mapping sheets can have blank separators; preserve coordinates.
                names = [v if v is not None else f"_blank_{i}" for i, v in enumerate(rows[0])]
            else:
                names = rows[0]
            frame = pd.DataFrame(rows[1:], columns=names)
            frame["_source_row"] = range(2, len(frame) + 2)
            sheets[ws.title] = frame
        book.close()
        return sheets
    except AuditError:
        raise
    except Exception as exc:
        raise AuditError("Unable to parse workbook; no fallback data were substituted.") from exc


def read_table(data, filename, sheet=None):
    if len(data) > MAX_BYTES:
        raise AuditError("File exceeds the 20 MB audit limit.")
    if data.startswith(b"PK\x03\x04"):
        books = read_workbook(data)
        if sheet is None and len(books) != 1:
            raise AuditError("Select an explicit worksheet for a multi-sheet workbook.")
        if sheet is not None and sheet not in books:
            raise AuditError("Requested worksheet is absent.")
        frame = books[sheet] if sheet else next(iter(books.values()))
        sheet = sheet or next(iter(books))
        frame["source_sheet"] = sheet
        actual_format = "OOXML workbook"
    else:
        try:
            text = data.decode("utf-8-sig")
            if not text.strip() or "\x00" in text:
                raise ValueError("not text")
            sep = "\t" if "\t" in text.splitlines()[0] else ","
            reader = csv.DictReader(StringIO(text), delimiter=sep)
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise ValueError("duplicate/missing headers")
            rows = []
            for row in reader:
                if None in row or any(v is None for v in row.values()):
                    raise ValueError("malformed row")
                row["_source_row"] = reader.line_num  # physical ending line, including quoted newlines
                rows.append(row)
            frame = pd.DataFrame(rows, columns=[*reader.fieldnames, "_source_row"] if "_source_row" not in reader.fieldnames else reader.fieldnames)
            actual_format = "TSV text" if sep == "\t" else "CSV text"
        except Exception as exc:
            raise AuditError("Expected UTF-8 CSV/TSV text or a macro-free XLSX workbook.") from exc
    frame["source_file"] = str(filename).replace("\\", "/").split("/")[-1]
    frame.attrs.update(actual_format=actual_format, sha256=hashlib.sha256(data).hexdigest(), sheet=sheet, source_file=frame["source_file"].iloc[0] if len(frame) else filename)
    return frame


def normalize_mortality(frame):
    frame = frame.rename(columns={"County": "county", "Month Code": "month", "Year": "year", "Population": "population"}).copy()
    if "county" in frame:
        frame["county"] = frame["county"].astype(str).str.replace(r" County, NY$", "", regex=True)
    return frame


def extract_document(data, filename):
    if len(data) > MAX_BYTES:
        raise AuditError("Evidence file exceeds the 20 MB audit limit.")
    name = str(filename).replace("\\", "/").split("/")[-1]
    units = []
    if data.startswith(b"%PDF"):
        from pypdf import PdfReader
        try:
            reader = PdfReader(BytesIO(data))
            for page, obj in enumerate(reader.pages, 1):
                text = obj.extract_text() or ""
                if not text.strip():
                    raise AuditError(f"No extractable text on page {page}; visual review is required.")
                units.append({"text": text, "page": page})
        except AuditError:
            raise
        except Exception as exc:
            raise AuditError("PDF extraction failed.") from exc
    elif data.startswith(b"PK\x03\x04"):
        import re
        import xml.etree.ElementTree as ET
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                    raise AuditError("Document expansion exceeds the 100 MB limit.")
                if any("vbaProject" in x for x in archive.namelist()):
                    raise AuditError("Macro-enabled documents are not supported.")
                if "word/document.xml" in archive.namelist():
                    root = ET.fromstring(archive.read("word/document.xml"))
                    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                    lines = ["".join(t.text or "" for t in p.iter(ns + "t")) for p in root.iter(ns + "p")]
                    units.append({"text": "\n".join(lines), "page": None,
                        "locator": "DOCX main-document paragraphs, including tables; no rendered page numbers"})
                else:
                    slides = sorted((x for x in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", x)),
                        key=lambda x: int(re.search(r"slide(\d+)", x).group(1)))
                    if not slides:raise AuditError("Expected a DOCX or PPTX document.")
                    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                    for path in slides:
                        root = ET.fromstring(archive.read(path))
                        text = "\n".join("".join(t.text or "" for t in p.iter(ns + "t")) for p in root.iter(ns + "p"))
                        slide = int(re.search(r"slide(\d+)", path).group(1))
                        if not text.strip():raise AuditError(f"No extractable text on slide {slide}; visual review is required.")
                        units.append({"text": text, "page": None, "slide": slide, "locator": f"PPTX slide {slide}, extracted text paragraphs; images and notes are not interpreted"})
                if not any(x['text'].strip() for x in units):raise AuditError("Document has no extractable text; visual review is required.")
        except AuditError:
            raise
        except Exception as exc:
            raise AuditError("Unable to extract DOCX/PPTX text.") from exc
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeError as exc:
            raise AuditError("Expected a PDF or UTF-8 evidence document.") from exc
        if not text.strip() or "\x00" in text:
            raise AuditError("Empty or unsupported evidence document.")
        units.append({"text": text, "page": None})
    result = []
    for unit in units:
        lines = unit["text"].splitlines()
        for start in range(0, len(lines), 16):
            text = "\n".join(lines[start:start + 16])
            if text.strip():
                result.append({"source_file": name, "page": unit["page"], "line_start": start + 1,
                    "line_end": min(start + 16, len(lines)), "text": text,
                    "sha256": hashlib.sha256(text.encode()).hexdigest(), "classification": "private_local",
                    "chunk_id": f"{hashlib.sha256(data).hexdigest()[:12]}:{unit.get('slide', unit['page'])}:{start + 1}",
                    **({"locator": unit["locator"], "slide": unit.get("slide")} if "locator" in unit else {})})
    return result
